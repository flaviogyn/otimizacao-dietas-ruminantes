#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador do modelo de Programação Linear (PL) de custo mínimo da dieta - lote lac_cow
IFG - Pós-Graduação - Modelagem e Otimização Aplicada - Fazenda GO_campo_dourado

Monta a função objetivo, as variáveis de decisão e as restrições NASEM usadas pelo
otimizador de ração do RuFaS (RUFAS/biophysical/animal/ration/ration_optimizer.py),
calculando os coeficientes com as próprias funções do RuFaS.

Fontes de dados:
1. Custos (R$/kg MS): 'Fazenda Campo Dourado/ingredientes_racoes_GO_campo_dourado.xlsx', coluna I.
2. Composição dos alimentos: 'Fazenda Campo Dourado/GO_campo_dourado/inputs_resolvidos/NASEM_Comp.csv'.
3. Exigências da vaca média do lote lac_cow: simulação base do RuFaS (60 dias, 28/10/2026).

Saídas (nesta pasta):
- modelo_pl_lac_cow.md : formulação completa, com coeficientes numéricos.
- modelo_pl_lac_cow.lp : o mesmo modelo no formato CPLEX LP (lido por GLPK, CBC, HiGHS, etc.).

Execução (a partir da raiz do repositório):
    RuFaS-dev/.venv/bin/python otimizacao/gerar_modelo_pl.py
"""

import os
import sys
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import linprog

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
sys.path.insert(0, os.path.join(BASE_DIR, 'RuFaS-dev'))

from RUFAS.biophysical.animal.nutrients.nutrition_supply_calculator import (  # noqa: E402
    FeedInRation,
    NutritionSupplyCalculator as NSC,
)
from RUFAS.biophysical.animal.animal_module_constants import AnimalModuleConstants as AMC  # noqa: E402
from RUFAS.data_structures.feed_storage_to_animal_connection import (  # noqa: E402
    FeedCategorization,
    FeedComponentType,
)

EXCEL_PATH = os.path.join(BASE_DIR, 'Fazenda Campo Dourado', 'ingredientes_racoes_GO_campo_dourado.xlsx')
NASEM_PATH = os.path.join(BASE_DIR, 'Fazenda Campo Dourado', 'GO_campo_dourado', 'inputs_resolvidos', 'NASEM_Comp.csv')
OUT_DIR = os.path.dirname(os.path.abspath(__file__))

# Alimentos candidatos para vacas em lactação. Ficam de fora:
#   23  farinha de sangue -> proibida para ruminantes (IN MAPA nº 8/2004)
#   202 leite integral e 216 starter -> alimentos exclusivos de bezerros
FEEDS = [18, 44, 51, 94, 100, 110, 176, 301, 303]

# Exigências médias por vaca do lote LAC_COW_PEN_3 (média das formulações dos dias 1 e 46 da
# simulação base). Variáveis: AnimalModuleReporter.report_average_nutrient_requirements.avg_rqmts_for_LAC_COW_PEN_3
REQ = {
    'DMI': 25.187,       # DMIest_requirement (kg MS/dia)
    'BW': 640.50,        # avg_BW (kg)
    'NE_maint': 12.686,  # NEmaint_requirement (Mcal)
    'NE_lact': 32.450,   # NEl_requirement (Mcal)
    'NE_growth': 0.702,  # NEg_requirement (Mcal)
    'NE_preg': 2.259,    # NEpreg_requirement (Mcal)
    'NE_act': 0.0,       # NEa_requirement (Mcal)
    'MP': 2742.94,       # MP_requirement (g)
    'Ca': 68.783,        # Ca_requirement (g)
    'P': 65.517,         # max(P_req, P_req_process) (g)
}
REQ['NE_total'] = REQ['NE_maint'] + REQ['NE_lact'] + REQ['NE_growth'] + REQ['NE_preg'] + REQ['NE_act']

# Perdas energéticas médias por vaca na simulação base (entram como constantes na energia NASEM).
CH4_G = 349.3        # enteric_methane_emission_for_LAC_COW_PEN_3 / nº de vacas (g/dia)
URINE_N_KG = 0.2795  # LAC_COW_PEN_3_urine_nitrogen / nº de vacas (kg/dia)

# Ração do caso base (kg MS/vaca/dia), usada como ponto de referência das linearizações.
REFERENCE_RATION = {44: 5.98, 51: 7.78, 58: 2.24, 110: 5.01, 167: 3.99, 301: 0.75}

FEED_NAMES_PT = {
    18: 'Feno de bermudagrass (Tifton)',
    44: 'Milho grão moído seco',
    51: 'Silagem de milho',
    94: 'Feno capim-leguminosa',
    100: 'Silagem capim-leguminosa',
    110: 'Silagem de leguminosa',
    176: 'Silagem de capim-sudão',
    301: 'Premix mineral/vitamínico',
    303: 'Blend de subprodutos (conc. proteico)',
}

MP_ENDOGENOUS = 0.4 * 11.8  # PM endógena do RuFaS, g por kg de MS consumida

CA_DIGESTIBILITY = {FeedComponentType.FORAGE: 0.3, FeedComponentType.CONC: 0.6, FeedComponentType.MINERAL: 0.95}
P_DIGESTIBILITY = {FeedComponentType.FORAGE: 0.64, FeedComponentType.CONC: 0.7, FeedComponentType.MINERAL: 0.8}


def load_feed_library() -> dict[int, SimpleNamespace]:
    library = pd.read_csv(NASEM_PATH).set_index('rufas_id')
    feeds = {}
    for rufas_id in set(FEEDS) | set(REFERENCE_RATION):
        row = library.loc[rufas_id].to_dict()
        row['rufas_id'] = rufas_id
        row['feed_type'] = FeedComponentType(row['feed_type'])
        row['Fd_Category'] = FeedCategorization(row['Fd_Category'])
        feeds[rufas_id] = SimpleNamespace(**row)
    return feeds


def load_costs() -> dict[int, tuple[float, float]]:
    ws = openpyxl.load_workbook(EXCEL_PATH).active
    costs = {}
    for r in range(4, ws.max_row + 1):
        fid = ws.cell(row=r, column=1).value
        if fid is None:
            continue
        price_ton = ws.cell(row=r, column=10).value
        costs[int(fid)] = (float(ws.cell(row=r, column=9).value), price_ton if isinstance(price_ton, (int, float)) else None)
    return costs


def reference_point(lib: dict[int, SimpleNamespace]) -> dict[str, float]:
    """Avalia, na ração do caso base, os termos não lineares que o RuFaS recalcula a cada formulação."""
    ref = [FeedInRation(amount, lib[fid]) for fid, amount in REFERENCE_RATION.items()]
    dmi = sum(REFERENCE_RATION.values())
    return {
        # O RuFaS passa o amido total em kg (não em % da MS) para calculate_NASEM_dNDF; mantido igual.
        'total_starch': sum(f.amount * f.info.starch / 100 for f in ref),
        'tdn_discount': NSC.calculate_nutrient_intake_discount(ref, REQ['BW']),
        'conc_pct': NSC._calculate_percentage_of_concentrates(ref, dmi),
    }


def build_coefficients(lib, costs, ref) -> pd.DataFrame:
    unit_feeds = [FeedInRation(1.0, lib[fid]) for fid in FEEDS]
    passage = NSC._calculate_protein_passage_rates(unit_feeds, REQ['DMI'], REQ['BW'], ref['conc_pct'])
    rdp_pct = NSC._calculate_rumen_degradable_protein_percentages(unit_feeds, passage)
    rup_pct = NSC._calculate_rumen_undegradable_protein_percentages(unit_feeds, rdp_pct)

    rows = []
    for f in unit_feeds:
        info = f.info
        de = NSC.calculate_NASEM_digestible_energy([f], REQ['DMI'], REQ['BW'], ref['total_starch'])
        tdn = info.TDN * ref['tdn_discount'] / 100
        rup_dig = rup_pct[info.rufas_id] / 100 * info.dRUP / 100 * 1000
        rows.append({
            'feed': info.rufas_id,
            'nome': FEED_NAMES_PT[info.rufas_id],
            'tipo': info.feed_type.value,
            'custo': costs[info.rufas_id][0],
            'preco_t_mn': costs[info.rufas_id][1],
            'DE': de,                                         # Mcal/kg MS (NASEM, RuFaS)
            'NE': AMC.EFF_OF_ME_USE * de,                     # Mcal/kg MS
            # g PM/kg MS = microbiana (limitada por NDT ou por PDR) + PNDR digestível + endógena (0,4 x 11,8 g/kg MS)
            'MP_tdn': 0.64 * 0.13 * 1000 * tdn + rup_dig + MP_ENDOGENOUS,
            'MP_rdp': 0.64 * 0.85 * 1000 * rdp_pct[info.rufas_id] / 100 + rup_dig + MP_ENDOGENOUS,
            'Ca': CA_DIGESTIBILITY.get(info.feed_type, 0.0) * info.calcium * 10,    # g Ca digestível/kg MS
            'P': P_DIGESTIBILITY.get(info.feed_type, 0.0) * info.phosphorus * 10,   # g P digestível/kg MS
            'NDF': info.NDF,
            'NDF_forr': info.NDF if info.feed_type is FeedComponentType.FORAGE else 0.0,
            'EE': info.EE,
            'limit': info.limit,
        })
    return pd.DataFrame(rows).set_index('feed')


def build_constraints(c: pd.DataFrame) -> tuple[list[dict], float]:
    """Cada restrição: nome, coeficientes por alimento, sentido, lado direito, descrição e origem no RuFaS."""
    D = REQ['DMI']
    f = AMC.DMI_CONSTRAINT_FRACTION
    energy_losses = 13.28 * CH4_G / 1000 + 0.0146 * URINE_N_KG * 1000
    ones = pd.Series(1.0, index=c.index)
    return [
        {'nome': 'R1a_consumo_MS_min', 'coef': ones, 'sentido': '>=', 'rhs': D * (1 - f),
         'desc': f'Consumo de MS ≥ (1 − {fmt(f, 2)}) × consumo estimado (kg MS/dia)', 'origem': 'DMI_constraint_lower'},
        {'nome': 'R1b_consumo_MS_max', 'coef': ones, 'sentido': '<=', 'rhs': D * (1 + f),
         'desc': f'Consumo de MS ≤ (1 + {fmt(f, 2)}) × consumo estimado (kg MS/dia)', 'origem': 'DMI_constraint_upper'},
        {'nome': 'R2_energia_liquida', 'coef': c['NE'], 'sentido': '>=',
         'rhs': REQ['NE_total'] + AMC.EFF_OF_ME_USE * energy_losses,
         'desc': 'Energia líquida NASEM ≥ manutenção + lactação + crescimento + gestação (Mcal/dia)',
         'origem': 'NASEM_net_energy_constraint'},
        {'nome': 'R3_PM_min_NDT', 'coef': c['MP_tdn'], 'sentido': '>=', 'rhs': REQ['MP'],
         'desc': 'Proteína metabolizável ≥ exigência, com proteína microbiana limitada pelo NDT (g/dia)',
         'origem': 'protein_constraint_lower (1º termo do min)'},
        {'nome': 'R4_PM_min_PDR', 'coef': c['MP_rdp'], 'sentido': '>=', 'rhs': REQ['MP'],
         'desc': 'Proteína metabolizável ≥ exigência, com proteína microbiana limitada pela PDR (g/dia)',
         'origem': 'protein_constraint_lower (2º termo do min)'},
        {'nome': 'R5_PM_max', 'coef': c['MP_tdn'], 'sentido': '<=',
         'rhs': AMC.PROTEIN_UPPER_LIMIT_FACTOR * REQ['MP'],
         'desc': f'Proteína metabolizável ≤ {fmt(AMC.PROTEIN_UPPER_LIMIT_FACTOR, 1)} × exigência (g/dia)',
         'origem': 'protein_constraint_upper'},
        {'nome': 'R6_calcio', 'coef': c['Ca'], 'sentido': '>=', 'rhs': REQ['Ca'],
         'desc': 'Cálcio digestível ≥ exigência (g/dia)', 'origem': 'calcium_constraint'},
        {'nome': 'R7_fosforo', 'coef': c['P'], 'sentido': '>=', 'rhs': REQ['P'],
         'desc': 'Fósforo digestível ≥ exigência (g/dia)', 'origem': 'phosphorus_constraint'},
        {'nome': 'R8_FDN_min', 'coef': c['NDF'] - AMC.MINIMUM_RATION_NDF, 'sentido': '>=', 'rhs': 0.0,
         'desc': f'FDN da dieta ≥ {AMC.MINIMUM_RATION_NDF:g}% da MS', 'origem': 'NDF_constraint_lower'},
        {'nome': 'R9_FDN_max', 'coef': c['NDF'] - AMC.MAXIMUM_RATION_NDF, 'sentido': '<=', 'rhs': 0.0,
         'desc': f'FDN da dieta ≤ {AMC.MAXIMUM_RATION_NDF:g}% da MS', 'origem': 'NDF_constraint_upper'},
        {'nome': 'R10_FDN_forragem', 'coef': c['NDF_forr'] - AMC.MINIMUM_RATION_FORAGE_NDF, 'sentido': '>=',
         'rhs': 0.0, 'desc': f'FDN de forragem ≥ {AMC.MINIMUM_RATION_FORAGE_NDF:g}% da MS',
         'origem': 'forage_NDF_constraint'},
        {'nome': 'R11_extrato_etereo', 'coef': c['EE'] - AMC.MAXIMUM_RATION_FAT, 'sentido': '<=', 'rhs': 0.0,
         'desc': f'Extrato etéreo da dieta ≤ {AMC.MAXIMUM_RATION_FAT:g}% da MS', 'origem': 'fat_constraint'},
    ], energy_losses


def fmt(v: float, nd: int = 4) -> str:
    return f'{v:,.{nd}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def linear_expr(coef: pd.Series, nd: int = 4) -> str:
    terms = []
    for fid, v in coef.items():
        if abs(v) < 1e-12:
            continue
        sign = '−' if v < 0 else '+'
        terms.append(f'{sign} {fmt(abs(v), nd)}·x{fid}')
    expr = ' '.join(terms)
    return expr[2:] if expr.startswith('+ ') else '−' + expr[2:]


def write_markdown(c, cons, ref, energy_losses, check) -> str:
    sense_md = {'>=': '≥', '<=': '≤', '=': '='}
    L = []
    L.append('# Modelo de PL – Dieta de custo mínimo das vacas em lactação (GO_campo_dourado)\n')
    L.append('> Gerado por `otimizacao/gerar_modelo_pl.py`. Não edite à mão: altere a planilha ou o script e gere de novo.\n')
    L.append('## 1. Escopo\n')
    L.append('- **Lote:** `lac_cow` (vacas em lactação), formulado para a **vaca média** do lote, base **1 dia**.')
    L.append('- **Padrão nutricional:** NASEM (`nutrient_standard` do `config.json`), com as mesmas restrições do '
             'otimizador de ração do RuFaS (`RUFAS/biophysical/animal/ration/ration_optimizer.py`).')
    L.append(f'- **Peso vivo médio:** {fmt(REQ["BW"], 1)} kg · **Consumo de MS estimado:** {fmt(REQ["DMI"], 3)} kg/dia.\n')

    L.append('## 2. Variáveis de decisão\n')
    L.append('xᵢ = quantidade do alimento *i* na dieta, em **kg de matéria seca por vaca por dia** (xᵢ ≥ 0).\n')
    L.append('| Variável | Alimento | Tipo (RuFaS) | Custo cᵢ (R$/kg MS) | Preço (R$/t MN) |')
    L.append('|---|---|---|---|---|')
    for fid, r in c.iterrows():
        t = fmt(r.preco_t_mn, 2) if r.preco_t_mn is not None else '—'
        L.append(f'| x{fid} | {r.nome} | {r.tipo} | {fmt(r.custo)} | {t} |')
    L.append('\nFicam **fora** do modelo do lote `lac_cow`:')
    L.append('- **x23 – farinha de sangue:** uso proibido para ruminantes no Brasil (IN MAPA nº 8/2004).')
    L.append('- **x202 – leite integral** e **x216 – starter:** alimentos exclusivos de bezerros (lote `calf`).\n')

    L.append('## 3. Função objetivo\n')
    L.append('Minimizar o custo diário da dieta por vaca (R$/vaca/dia):\n')
    L.append('$$\\min Z = \\sum_i c_i \\, x_i$$\n')
    L.append('```')
    L.append(f'min Z = {linear_expr(c["custo"])}')
    L.append('```\n')

    L.append('## 4. Restrições\n')
    for k in cons:
        L.append(f'**{k["nome"]}** – {k["desc"]}  ')
        L.append(f'*Origem no RuFaS:* `{k["origem"]}`\n')
        L.append('```')
        L.append(f'{linear_expr(k["coef"])}  {sense_md[k["sentido"]]}  {fmt(k["rhs"])}')
        L.append('```\n')
    L.append('**Não negatividade:** xᵢ ≥ 0 para todo *i*. O limite superior do `NASEM_Comp.csv` '
             f'(`limit` = {fmt(c["limit"].max(), 0)} kg) nunca é atingido, porque Σxᵢ ≤ {fmt(cons[1]["rhs"], 3)} kg.\n')

    L.append('## 5. Matriz de coeficientes (para o Solver)\n')
    head = '| Restrição | ' + ' | '.join(f'x{f}' for f in c.index) + ' | Sentido | Lado direito |'
    L.append(head)
    L.append('|' + '---|' * (len(c.index) + 3))
    L.append('| **Custo (Z)** | ' + ' | '.join(fmt(v) for v in c['custo']) + ' | min | — |')
    for k in cons:
        L.append(f'| {k["nome"]} | ' + ' | '.join(fmt(v) for v in k['coef']) +
                 f' | {sense_md[k["sentido"]]} | {fmt(k["rhs"])} |')
    L.append('')

    L.append('## 6. De onde vem cada coeficiente\n')
    L.append('| Coeficiente | Cálculo | Unidade |')
    L.append('|---|---|---|')
    L.append('| cᵢ | coluna I da planilha (`purchased_feed_cost`) | R$/kg MS |')
    L.append(f'| NEᵢ (R2) | {fmt(AMC.EFF_OF_ME_USE, 2)} × EDᵢ NASEM, EDᵢ por `calculate_NASEM_digestible_energy` | Mcal/kg MS |')
    L.append('| PM (R3, R5) | 0,64 × 0,13 × 1000 × NDTᵢ × desconto + PNDRᵢ × dPNDRᵢ × 1000 + 0,4 × 11,8 (endógena) | g/kg MS |')
    L.append('| PM (R4) | 0,64 × 0,85 × 1000 × PDRᵢ + PNDRᵢ × dPNDRᵢ × 1000 + 0,4 × 11,8 (endógena) | g/kg MS |')
    L.append('| Caᵢ (R6) | digestibilidade × Ca%ᵢ × 10 (forragem 0,30; concentrado 0,60; mineral 0,95) | g/kg MS |')
    L.append('| Pᵢ (R7) | digestibilidade × P%ᵢ × 10 (forragem 0,64; concentrado 0,70; mineral 0,80) | g/kg MS |')
    L.append('| FDN, FDN de forragem, EE (R8–R11) | teor do alimento − limite em % (só forragens contam em R10) | % MS |')
    L.append('')
    L.append('**Lados direitos:**\n')
    L.append(f'- R2: exigência de energia líquida = {fmt(REQ["NE_maint"], 3)} (manutenção) + {fmt(REQ["NE_lact"], 3)} '
             f'(lactação) + {fmt(REQ["NE_growth"], 3)} (crescimento) + {fmt(REQ["NE_preg"], 3)} (gestação) = '
             f'**{fmt(REQ["NE_total"], 3)} Mcal**. Como o RuFaS desconta da energia digestível as perdas por metano '
             f'(13,28 × {fmt(CH4_G / 1000, 4)} kg CH₄) e por urina (0,0146 × {fmt(URINE_N_KG * 1000, 1)} g N) = '
             f'{fmt(energy_losses, 3)} Mcal, esse termo constante passa para o lado direito: '
             f'{fmt(REQ["NE_total"], 3)} + {fmt(AMC.EFF_OF_ME_USE, 2)} × {fmt(energy_losses, 3)}.')
    L.append(f'- R1a/R1b: consumo estimado de {fmt(REQ["DMI"], 3)} kg ± {AMC.DMI_CONSTRAINT_FRACTION:.0%}.')
    L.append(f'- R3–R5: exigência de PM = {fmt(REQ["MP"], 2)} g; limite superior = '
             f'{fmt(AMC.PROTEIN_UPPER_LIMIT_FACTOR, 1)} × exigência. A PM endógena do RuFaS (0,4 × 11,8 g por kg de MS) '
             'é proporcional ao consumo e por isso entra nos coeficientes, não no lado direito.')
    L.append(f'- R6, R7: Ca = {fmt(REQ["Ca"], 3)} g e P = {fmt(REQ["P"], 3)} g (o maior entre `P_req` e `P_req_process`).\n')

    L.append('## 7. Hipóteses e linearizações\n')
    L.append('O otimizador do RuFaS é **não linear** (SLSQP). Para virar um PL resolvível pelo Simplex:')
    L.append(f'1. **Restrições em % da MS reescritas na forma linear:** por exemplo, FDN ≥ 25% vira '
             'Σ(FDNᵢ − 25)·xᵢ ≥ 0. O consumo de MS pode variar na mesma faixa do RuFaS '
             f'(±{AMC.DMI_CONSTRAINT_FRACTION:.0%}, R1a/R1b); as digestibilidades que dependem do consumo '
             f'foram calculadas no consumo estimado ({fmt(REQ["DMI"], 3)} kg).')
    L.append('2. **Ponto de referência = ração do caso base** (simulação de 60 dias). Os termos que dependem da '
             'dieta inteira foram calculados nesse ponto e tratados como constantes:')
    L.append(f'   - desconto de NDT por nível de consumo = {fmt(ref["tdn_discount"], 4)};')
    L.append(f'   - amido total da dieta (entra na digestibilidade da FDN) = {fmt(ref["total_starch"], 3)} kg;')
    L.append(f'   - concentrado na dieta (entra na taxa de passagem da proteína) = {fmt(ref["conc_pct"], 2)}%.')
    L.append('3. **PM com mínimo de dois termos:** a proteína microbiana do RuFaS é min(limite por NDT, limite por PDR). '
             'Exigir PM ≥ exigência com o mínimo equivale a exigir as duas restrições lineares R3 **e** R4. '
             'Para o limite superior (R5) foi usado o termo limitado pelo NDT, o que é uma condição suficiente.')
    L.append('4. **Exigências fixas** da vaca média: média das duas formulações do caso base (dias 1 e 46). '
             'Perdas de metano e N urinário também vêm do caso base.')
    L.append('5. **Observação sobre o RuFaS:** `calculate_NASEM_dNDF` usa `(total_starch − 26)`; a equação NASEM '
             'espera o amido em % da MS, mas o RuFaS passa em kg. O modelo reproduz o cálculo do RuFaS para que '
             'a solução seja coerente com o simulador.\n')

    L.append('## 8. Verificação\n')
    L.append(f'- O modelo é **{check["status"]}** (resolvido com HiGHS via `scipy.optimize.linprog`).')
    if check['ok']:
        L.append(f'- Na solução do PL, recalculando os termos não lineares com as funções do RuFaS, '
                 f'a energia líquida fica {fmt(check["ne_err"], 2)}% e a PM {fmt(check["mp_err"], 2)}% '
                 'em relação ao valor do PL (erro da linearização).')
    L.append('')
    return '\n'.join(L)


def write_lp(c, cons) -> str:
    def expr(coef):
        out = []
        for fid, v in coef.items():
            if abs(v) < 1e-12:
                continue
            out.append(f'{"-" if v < 0 else "+"} {abs(v):.6f} x{fid}')
        s = ' '.join(out)
        return s[2:] if s.startswith('+ ') else s
    L = ['\\ Dieta de custo minimo - lote lac_cow - GO_campo_dourado (gerado por gerar_modelo_pl.py)',
         '\\ x<id> = kg de materia seca por vaca por dia do alimento <id>', 'Minimize',
         f' custo: {expr(c["custo"])}', 'Subject To']
    for k in cons:
        L.append(f' {k["nome"]}: {expr(k["coef"])} {k["sentido"]} {k["rhs"]:.6f}')
    L.append('Bounds')
    L += [f' x{fid} >= 0' for fid in c.index]
    L.append('End')
    return '\n'.join(L) + '\n'


def check_model(c, cons, lib, ref) -> dict:
    A_ub, b_ub, A_eq, b_eq = [], [], [], []
    for k in cons:
        if k['sentido'] == '=':
            A_eq.append(k['coef'].values); b_eq.append(k['rhs'])
        elif k['sentido'] == '<=':
            A_ub.append(k['coef'].values); b_ub.append(k['rhs'])
        else:
            A_ub.append(-k['coef'].values); b_ub.append(-k['rhs'])
    res = linprog(c['custo'].values, A_ub=np.array(A_ub), b_ub=b_ub, A_eq=np.array(A_eq) if A_eq else None,
                  b_eq=b_eq or None,
                  bounds=[(0, None)] * len(c), method='highs')
    if res.status != 0:
        return {'ok': False, 'status': f'inviável/sem solução ({res.message})', 'res': res}

    # Avalia energia e PM com as equações não lineares do RuFaS na solução do PL.
    x = dict(zip(c.index, res.x))
    feeds = [FeedInRation(v, lib[fid]) for fid, v in x.items() if v > 1e-9]
    dmi = sum(x.values())
    starch = sum(f.amount * f.info.starch / 100 for f in feeds)
    me = NSC.calculate_NASEM_metabolizable_energy(feeds, dmi, REQ['BW'], starch, CH4_G, URINE_N_KG)
    ne_true = NSC.calculate_NASEM_net_energy(me)
    disc = NSC.calculate_nutrient_intake_discount(feeds, REQ['BW'])
    mp_true = NSC.calculate_metabolizable_protein_supply(
        feeds, dmi, {f.info.rufas_id: f.info.TDN * disc for f in feeds}, REQ['BW'])
    ne_lp = float(c['NE'] @ res.x) - AMC.EFF_OF_ME_USE * (13.28 * CH4_G / 1000 + 0.0146 * URINE_N_KG * 1000)
    mp_lp = min(float(c['MP_tdn'] @ res.x), float(c['MP_rdp'] @ res.x))
    return {'ok': True, 'status': 'viável', 'res': res, 'x': x,
            'ne_true': ne_true, 'ne_lp': ne_lp, 'ne_err': (ne_true / ne_lp - 1) * 100,
            'mp_true': mp_true, 'mp_lp': mp_lp, 'mp_err': (mp_true / mp_lp - 1) * 100}


def main():
    lib = load_feed_library()
    costs = load_costs()
    ref = reference_point(lib)
    c = build_coefficients(lib, costs, ref)
    cons, energy_losses = build_constraints(c)
    check = check_model(c, cons, lib, ref)

    md_path = os.path.join(OUT_DIR, 'modelo_pl_lac_cow.md')
    lp_path = os.path.join(OUT_DIR, 'modelo_pl_lac_cow.lp')
    with open(md_path, 'w', encoding='utf-8') as fh:
        fh.write(write_markdown(c, cons, ref, energy_losses, check))
    with open(lp_path, 'w', encoding='utf-8') as fh:
        fh.write(write_lp(c, cons))

    print(f'Modelo gravado em:\n  {md_path}\n  {lp_path}')
    print(f'Status: {check["status"]}')
    if check['ok']:
        print(f'Custo ótimo: R$ {check["res"].fun:.4f}/vaca/dia')
        for fid, v in check['x'].items():
            if v > 1e-6:
                print(f'  x{fid:<4d} {FEED_NAMES_PT[fid]:32s} {v:8.3f} kg MS ({v / check["res"].x.sum() * 100:5.1f}%)')
        print(f'NE  PL={check["ne_lp"]:.2f}  RuFaS={check["ne_true"]:.2f} Mcal ({check["ne_err"]:+.2f}%)')
        print(f'PM  PL={check["mp_lp"]:.1f}  RuFaS={check["mp_true"]:.1f} g ({check["mp_err"]:+.2f}%)')


if __name__ == '__main__':
    main()
