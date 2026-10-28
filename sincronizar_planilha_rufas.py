#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de Sincronização: Planilha de Ingredientes -> Inputs do RuFaS
IFG - Mestrado / Pós-Graduação em IA - Otimização de Dietas de Ruminantes

Este script:
1. Lê 'Fazenda Campo Dourado/ingredientes_racoes_GO_campo_dourado.xlsx'.
2. Sincroniza os teores bromatológicos no 'NASEM_Comp.csv' (DM, CP, NDF, Ca, P).
3. Sincroniza os custos de compra (purchased_feed_cost) no 'feed.json'.
4. Propaga os arquivos atualizados para 'RuFaS-dev/GO_campo_dourado/inputs_resolvidos/'.
"""

import os
import json
import shutil
import openpyxl
import pandas as pd

def sync_excel_to_rufas():
    base_dir = os.path.abspath(os.path.dirname(__file__))
    
    excel_path = os.path.join(base_dir, 'Fazenda Campo Dourado', 'ingredientes_racoes_GO_campo_dourado.xlsx')
    fazenda_inputs_dir = os.path.join(base_dir, 'Fazenda Campo Dourado', 'GO_campo_dourado', 'inputs_resolvidos')
    rufas_inputs_dir = os.path.join(base_dir, 'RuFaS-dev', 'GO_campo_dourado', 'inputs_resolvidos')
    
    nasem_csv_path = os.path.join(fazenda_inputs_dir, 'NASEM_Comp.csv')
    feed_json_path = os.path.join(fazenda_inputs_dir, 'feed.json')

    print("=" * 70)
    print("SINCRONIZADOR: PLANILHA EXCEL -> INPUTS RUFAS (GO_CAMPO_DOURADO)")
    print("=" * 70)

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Planilha não encontrada: {excel_path}")
    if not os.path.exists(nasem_csv_path):
        raise FileNotFoundError(f"NASEM_Comp.csv não encontrado: {nasem_csv_path}")
    if not os.path.exists(feed_json_path):
        raise FileNotFoundError(f"feed.json não encontrado: {feed_json_path}")

    # 1. Carregar planilha
    print(f"\n[1/4] Lendo dados da planilha Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    feeds_from_excel = {}
    for r in range(4, ws.max_row + 1):
        fid_val = ws.cell(row=r, column=1).value
        if fid_val is None:
            continue
        fid = int(fid_val)
        nome = str(ws.cell(row=r, column=2).value or '').strip()
        cat = str(ws.cell(row=r, column=3).value or '').strip()
        dm = float(ws.cell(row=r, column=4).value or 0.0)
        cp = float(ws.cell(row=r, column=5).value or 0.0)
        ndf = float(ws.cell(row=r, column=6).value or 0.0)
        ca = float(ws.cell(row=r, column=7).value or 0.0)
        p = float(ws.cell(row=r, column=8).value or 0.0)
        custo = float(ws.cell(row=r, column=9).value or 0.0)

        feeds_from_excel[fid] = {
            'nome': nome,
            'categoria': cat,
            'DM': dm,
            'CP': cp,
            'NDF': ndf,
            'calcium': ca,
            'phosphorus': p,
            'cost': custo
        }
        print(f"  - feed_type {fid:3d} ({nome[:28]:28s}): DM={dm:5.2f}% | CP={cp:5.2f}% | NDF={ndf:5.2f}% | Ca={ca:5.3f}% | P={p:5.3f}% | R$/kg={custo:.2f}")

    # 2. Atualizar NASEM_Comp.csv
    print(f"\n[2/4] Atualizando matriz bromatológica em: {nasem_csv_path}")
    df_nasem = pd.read_csv(nasem_csv_path)

    updated_count = 0
    for fid, vals in feeds_from_excel.items():
        mask = df_nasem['rufas_id'] == fid
        if mask.any():
            df_nasem.loc[mask, 'DM'] = vals['DM']
            df_nasem.loc[mask, 'CP'] = vals['CP']
            df_nasem.loc[mask, 'NDF'] = vals['NDF']
            df_nasem.loc[mask, 'calcium'] = vals['calcium']
            df_nasem.loc[mask, 'phosphorus'] = vals['phosphorus']
            updated_count += 1
        else:
            print(f"  [AVISO] rufas_id {fid} não encontrado no NASEM_Comp.csv")

    df_nasem.to_csv(nasem_csv_path, index=False)
    print(f"  -> {updated_count} ingredientes atualizados com sucesso no NASEM_Comp.csv.")

    # 3. Atualizar feed.json com os novos custos
    print(f"\n[3/4] Atualizando custos no feed.json: {feed_json_path}")
    with open(feed_json_path, 'r', encoding='utf-8') as f:
        feed_config = json.load(f)

    cost_updates = 0
    for feed_item in feed_config.get('feeds', []):
        fid = feed_item.get('feed_type')
        if fid in feeds_from_excel:
            old_cost = feed_item.get('purchased_feed_cost')
            new_cost = feeds_from_excel[fid]['cost']
            feed_item['purchased_feed_cost'] = new_cost
            cost_updates += 1
            print(f"  - feed_type {fid}: custo atualizado de {old_cost} -> {new_cost} R$/kg MS")

    with open(feed_json_path, 'w', encoding='utf-8') as f:
        json.dump(feed_config, f, indent=2)
    print(f"  -> {cost_updates} custos atualizados no feed.json.")

    # 4. Sincronizar com RuFaS-dev/GO_campo_dourado/inputs_resolvidos/
    print(f"\n[4/4] Propagando arquivos para o simulador ativo em: {rufas_inputs_dir}")
    if os.path.exists(rufas_inputs_dir):
        shutil.copy2(nasem_csv_path, os.path.join(rufas_inputs_dir, 'NASEM_Comp.csv'))
        shutil.copy2(feed_json_path, os.path.join(rufas_inputs_dir, 'feed.json'))
        print("  -> NASEM_Comp.csv e feed.json copiados com sucesso para RuFaS-dev/GO_campo_dourado!")
    else:
        print(f"  [AVISO] Pasta {rufas_inputs_dir} não encontrada. Certifique-se de copiar o cenário para RuFaS-dev.")

    print("\n" + "=" * 70)
    print("SINCRONIZAÇÃO CONCLUÍDA COM SUCESSO! O RUFAS ESTÁ PRONTO PARA SIMULAR.")
    print("=" * 70)

if __name__ == '__main__':
    sync_excel_to_rufas()
