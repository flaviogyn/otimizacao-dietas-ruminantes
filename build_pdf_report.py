import os
from xml.sax.saxutils import escape

import openpyxl
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

EXCEL_PATH = os.path.join('Fazenda Campo Dourado', 'ingredientes_racoes_GO_campo_dourado.xlsx')


def load_spreadsheet():
    """Lê as linhas de alimentos da planilha (cabeçalho na linha 3, dados a partir da linha 4)."""
    ws = openpyxl.load_workbook(EXCEL_PATH).active
    rows = []
    for r in range(4, ws.max_row + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, 14)]
        if values[0] is None:
            continue
        rows.append({
            'feed_type': int(values[0]), 'nome': values[1], 'categoria': values[2],
            'DM': values[3], 'CP': values[4], 'NDF': values[5], 'Ca': values[6], 'P': values[7],
            'custo': values[8], 'preco_t': values[9],
            'preco_original': values[10] or '—', 'fonte': values[11] or '—', 'conversao': values[12] or '—',
        })
    return rows


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_decorations(self, page_count):
        self.saveState()
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.setFont('Helvetica', 8)
            self.setFillColor(colors.HexColor('#64748b'))
            self.drawString(36, letter[1] - 25, 'IFG / RuFaS — Fazenda Campo Dourado | Relatório de Composição Bromatológica e Preços')
            self.drawRightString(letter[0] - 36, letter[1] - 25, 'Outubro / 2026')
            self.setStrokeColor(colors.HexColor('#cbd5e1'))
            self.setLineWidth(0.5)
            self.line(36, letter[1] - 28, letter[0] - 36, letter[1] - 28)

        # Footer (all pages)
        self.setFont('Helvetica', 8)
        self.setFillColor(colors.HexColor('#64748b'))
        page_text = f'Página {self._pageNumber} de {page_count}'
        self.drawRightString(letter[0] - 36, 22, page_text)
        self.drawString(36, 22, 'Confidencial / Uso Acadêmico — IFG — Otimização de Dietas de Ruminantes')
        self.setStrokeColor(colors.HexColor('#cbd5e1'))
        self.setLineWidth(0.5)
        self.line(36, 32, letter[0] - 36, 32)
        self.restoreState()


def generate_pdf():
    pdf_filename = r'Fazenda Campo Dourado/relatorio_ingredientes_racoes_GO_campo_dourado.pdf'
    docs_pdf_filename = r'docs/relatorio_ingredientes_racoes_GO_campo_dourado.pdf'
    
    os.makedirs('Fazenda Campo Dourado', exist_ok=True)
    os.makedirs('docs', exist_ok=True)

    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=42
    )

    usable_width = letter[0] - 72  # 540 pt

    styles = getSampleStyleSheet()

    # Custom styles
    styles.add(ParagraphStyle(
        'MainTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=17,
        leading=21,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=4
    ))

    styles.add(ParagraphStyle(
        'SubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor('#1e40af'),
        spaceAfter=8
    ))

    styles.add(ParagraphStyle(
        'MetaBox',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#334155')
    ))

    styles.add(ParagraphStyle(
        'MetaBoxBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#0f172a')
    ))

    styles.add(ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    ))

    styles.add(ParagraphStyle(
        'BodyP',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    ))

    styles.add(ParagraphStyle(
        'BodyPBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=4
    ))

    styles.add(ParagraphStyle(
        'TableHead',
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=1 # Center
    ))

    styles.add(ParagraphStyle(
        'TableCell',
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1e293b')
    ))

    styles.add(ParagraphStyle(
        'TableCellCenter',
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1e293b'),
        alignment=1 # Center
    ))

    styles.add(ParagraphStyle(
        'TableCellRight',
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1e293b'),
        alignment=2 # Right
    ))

    styles.add(ParagraphStyle(
        'CalloutText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11.5,
        textColor=colors.HexColor('#0f766e')
    ))

    feeds = load_spreadsheet()
    story = []

    # 1. Header Banner Box
    banner_data = [
        [
            Paragraph("<b>IFG — INSTITUTO FEDERAL DE GOIÁS</b><br/>"
                      "Mestrado / Pós-Graduação em Inteligência Artificial & Zootecnia Computacional", styles['MetaBox']),
            Paragraph("<b>Projeto:</b> Otimização de Dietas de Ruminantes<br/>"
                      "<b>Simulador:</b> RuFaS (Ruminant Farm Systems)<br/>"
                      "<b>Padrão Nutricional:</b> NASEM (2021)", styles['MetaBox'])
        ]
    ]
    t_banner = Table(banner_data, colWidths=[usable_width*0.58, usable_width*0.42])
    t_banner.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_banner)
    story.append(Spacer(1, 10))

    # Title & Subtitle
    story.append(Paragraph("Relatório Técnico: Composição Bromatológica, Preços e Origem dos Dados", styles['MainTitle']))
    story.append(Paragraph("Consolidação dos 12 Alimentos da Fazenda Campo Dourado (GO_campo_dourado) com Base no RuFaS", styles['SubTitle']))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284c7'), spaceAfter=8, spaceBefore=0))

    # Meta Info Table
    meta_data = [
        [
            Paragraph("<b>Planilha Alvo:</b>", styles['TableCell']),
            Paragraph("ingredientes_racoes_GO_campo_dourado.xlsx", styles['TableCell']),
            Paragraph("<b>Cenário:</b>", styles['TableCell']),
            Paragraph("Conservador (100 vacas, Goiás)", styles['TableCell'])
        ],
        [
            Paragraph("<b>Arquivo-Mestre:</b>", styles['TableCell']),
            Paragraph("NASEM_Comp.csv (inputs_resolvidos)", styles['TableCell']),
            Paragraph("<b>Data do Ajuste:</b>", styles['TableCell']),
            Paragraph("Outubro / 2026", styles['TableCell'])
        ]
    ]
    t_meta = Table(meta_data, colWidths=[usable_width*0.17, usable_width*0.43, usable_width*0.16, usable_width*0.24])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f1f5f9')),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0,0), (-1,-1), 4),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 10))

    # Section 1: Overview
    story.append(Paragraph("1. Contexto e Objetivo do Documento", styles['SectionHeader']))
    story.append(Paragraph(
        "Este relatório formaliza a proveniência e os valores dos parâmetros nutricionais e dos preços inseridos na planilha "
        "<b><code>ingredientes_racoes_GO_campo_dourado.xlsx</code></b>. Os ingredientes correspondem aos 12 alimentos "
        "utilizados na formulação das dietas dos lotes de animais (bezerros, novilhas em recria, vacas no pré-parto e vacas em lactação) "
        "para o cenário representativo da bacia leiteira de Goiás no simulador biofísico <b>RuFaS</b>.",
        styles['BodyP']
    ))
    story.append(Paragraph(
        "A biblioteca de composição de alimentos adotada é a do padrão oficial <b>NASEM (2021)</b>, cujas características bromatológicas "
        "alimentam o motor de cálculo de exigências nutricionais e formulação linearizada de mínimo custo.",
        styles['BodyP']
    ))

    # Section 2: Main Consolidated Table
    story.append(Paragraph("2. Tabela de Alimentos Preenchida na Planilha", styles['SectionHeader']))
    story.append(Paragraph(
        "A tabela abaixo é lida diretamente da planilha Excel. A composição vem do NASEM (RuFaS); o custo "
        "(<code>purchased_feed_cost</code>, em R$/kg de MS) é o valor enviado ao RuFaS e usado no modelo de PL, e o "
        "preço por tonelada de matéria natural (MN) é a referência comercial:",
        styles['BodyP']
    ))

    # Larguras: 36 + 128 + 72 + 5 x 40 + 48 + 56 = 540 pt
    headers_main = [
        Paragraph("<b>feed_<br/>type</b>", styles['TableHead']),
        Paragraph("<b>Alimento</b>", styles['TableHead']),
        Paragraph("<b>Categoria</b>", styles['TableHead']),
        Paragraph("<b>DM<br/>(%MS)</b>", styles['TableHead']),
        Paragraph("<b>CP<br/>(%)</b>", styles['TableHead']),
        Paragraph("<b>NDF<br/>(%)</b>", styles['TableHead']),
        Paragraph("<b>Ca<br/>(%)</b>", styles['TableHead']),
        Paragraph("<b>P<br/>(%)</b>", styles['TableHead']),
        Paragraph("<b>Custo<br/>(R$/kg MS)</b>", styles['TableHead']),
        Paragraph("<b>Preço<br/>(R$/t MN)</b>", styles['TableHead'])
    ]

    rows_main = [headers_main]
    for item in feeds:
        preco_t = f"{item['preco_t']:,.2f}" if isinstance(item['preco_t'], (int, float)) else '—'
        rows_main.append([
            Paragraph(str(item['feed_type']), styles['TableCellCenter']),
            Paragraph(escape(item['nome']), styles['TableCell']),
            Paragraph(escape(item['categoria']), styles['TableCell']),
            Paragraph(f"{item['DM']:.2f}", styles['TableCellRight']),
            Paragraph(f"{item['CP']:.2f}", styles['TableCellRight']),
            Paragraph(f"{item['NDF']:.2f}", styles['TableCellRight']),
            Paragraph(f"{item['Ca']:.3f}", styles['TableCellRight']),
            Paragraph(f"{item['P']:.3f}", styles['TableCellRight']),
            Paragraph(f"{item['custo']:.4f}", styles['TableCellRight']),
            Paragraph(preco_t, styles['TableCellRight']),
        ])

    t_main = Table(rows_main, colWidths=[36, 128, 72, 40, 40, 40, 40, 40, 48, 56])
    
    t_main_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')), # Teal Header
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 3),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#0f766e')),
    ]
    # Alternating row colors
    for r_idx in range(1, len(rows_main)):
        if r_idx % 2 == 0:
            t_main_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#f8fafc')))
        else:
            t_main_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.white))
            
    t_main.setStyle(TableStyle(t_main_style))
    story.append(t_main)
    story.append(Spacer(1, 8))

    # Callout box regarding Ca and P
    callout_data = [[
        Paragraph(
            "<b>Nota Técnica sobre Cálcio e Fósforo (Linhas 4 a 6):</b><br/>"
            "Na planilha original, os alimentos 18, 23 e 44 haviam sido preenchidos com um deslocamento de 1 coluna "
            "(o teor de Fósforo constava sob <code>Ca (%)</code> e o de Magnésio sob <code>P (%)</code>). "
            "Todas as 12 linhas foram devidamente alinhadas e corrigidas conforme as variáveis <code>calcium</code> e "
            "<code>phosphorus</code> da matriz oficial do NASEM.",
            styles['CalloutText']
        )
    ]]
    t_callout = Table(callout_data, colWidths=[usable_width])
    t_callout.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f0fdfa')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#0d9488')),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_callout)
    story.append(Spacer(1, 6))

    # Callout box regarding feeds 23 and 303
    callout_feeds = [[
        Paragraph(
            "<b>Atenção aos alimentos 23 e 303:</b><br/>"
            "<b>23 – Farinha de sangue:</b> o uso de proteínas de origem animal na alimentação de ruminantes é proibido no "
            "Brasil (IN MAPA nº 8/2004). O custo de 999 R$/kg MS é uma penalidade para excluí-la da otimização.<br/>"
            "<b>303 – Blend de subprodutos:</b> no <code>NASEM_Comp.csv</code> ele é o <i>Farm ES Northeast BP Blend</i>, "
            "um concentrado (<code>feed_type = Conc</code>) com 33% de PB, e não um premix mineral, como constava na "
            "versão anterior da planilha. Foi precificado pelo DDGS de milho, o subproduto proteico equivalente "
            "no Centro-Oeste.",
            styles['CalloutText']
        )
    ]]
    t_callout_feeds = Table(callout_feeds, colWidths=[usable_width])
    t_callout_feeds.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#fff7ed')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#ea580c')),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_callout_feeds)
    story.append(Spacer(1, 14))

    # Section 3: Prices and sources
    story.append(Paragraph("3. Preços dos Alimentos e Fontes", styles['SectionHeader']))
    story.append(Paragraph(
        "Pesquisa de preços realizada em 28/10/2026. As cotações oficiais (CEPEA e CNA/Cepea) foram usadas quando "
        "existiam; para os demais itens foram usados anúncios, compras públicas e guias de preço, com as aproximações "
        "indicadas. Todos os preços foram convertidos para R$/kg de MS pela fórmula "
        "<b>R$/kg MS = (R$/kg MN) ÷ (MS% ÷ 100)</b>.",
        styles['BodyP']
    ))

    price_rows = [[
        Paragraph("<b>feed_<br/>type</b>", styles['TableHead']),
        Paragraph("<b>Preço original</b>", styles['TableHead']),
        Paragraph("<b>Fonte</b>", styles['TableHead']),
        Paragraph("<b>Conversão / observação</b>", styles['TableHead']),
    ]]
    for item in feeds:
        price_rows.append([
            Paragraph(str(item['feed_type']), styles['TableCellCenter']),
            Paragraph(escape(str(item['preco_original'])), styles['TableCell']),
            Paragraph(escape(str(item['fonte'])), styles['TableCell']),
            Paragraph(escape(str(item['conversao'])), styles['TableCell']),
        ])
    t_prices = Table(price_rows, colWidths=[36, 96, 214, 194], repeatRows=1)
    t_prices_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f766e')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 3),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#0f766e')),
    ]
    for r_idx in range(2, len(price_rows), 2):
        t_prices_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#f8fafc')))
    t_prices.setStyle(TableStyle(t_prices_style))
    story.append(t_prices)
    story.append(Spacer(1, 14))

    # Section 4: Data Sources
    story.append(Paragraph("4. Origem dos Dados e Rastreabilidade nos Arquivos do Projeto", styles['SectionHeader']))
    story.append(Paragraph(
        "Os teores nutricionais foram extraídos da estrutura de simulação da Fazenda Campo Dourado e os preços, da "
        "pesquisa da Seção 3. A cadeia de proveniência envolve os seguintes arquivos do repositório:",
        styles['BodyP']
    ))

    sources_data = [
        [
            Paragraph("<b>Arquivo / Caminho</b>", styles['TableHead']),
            Paragraph("<b>Escopo e Descrição no RuFaS</b>", styles['TableHead']),
            Paragraph("<b>Variáveis Extraídas</b>", styles['TableHead'])
        ],
        [
            Paragraph("<code>Fazenda Campo Dourado/<br/>GO_campo_dourado/<br/>inputs_resolvidos/<br/>NASEM_Comp.csv</code>", styles['TableCell']),
            Paragraph("Matriz bromatológica oficial com 176 alimentos do padrão NASEM (2021) e 90 variáveis por alimento. Idêntico ao arquivo na raiz do <code>RuFaS-dev</code>.", styles['TableCell']),
            Paragraph("<code>DM</code> (Matéria Seca),<br/><code>CP</code> (Proteína Bruta),<br/><code>NDF</code> (FDN),<br/><code>calcium</code> (Cálcio),<br/><code>phosphorus</code> (Fósforo).", styles['TableCell'])
        ],
        [
            Paragraph("<code>inputs_resolvidos/<br/>user_feeds.csv</code>", styles['TableCell']),
            Paragraph("Tabela descritiva de catálogo com 300 ingredientes, indexada por <code>rufas_id</code>.", styles['TableCell']),
            Paragraph("Nomes em inglês, descrição do alimento e categorias originais.", styles['TableCell'])
        ],
        [
            Paragraph("<code>inputs_resolvidos/<br/>feed.json</code> e<br/><code>feed_goias.json</code>", styles['TableCell']),
            Paragraph("Configuração das dietas dos lotes da fazenda e custos de aquisição por kg de Matéria Seca. "
                      "Recebe os custos da planilha via <code>sincronizar_planilha_rufas.py</code>.", styles['TableCell']),
            Paragraph("<code>purchased_feed_cost</code> (custo de compra por kg de MS).", styles['TableCell'])
        ],
        [
            Paragraph("<code>otimizacao/<br/>gerar_modelo_pl.py</code> e<br/><code>modelo_pl_lac_cow.md</code>", styles['TableCell']),
            Paragraph("Modelo de PL de custo mínimo do lote de vacas em lactação, com coeficientes calculados pelas "
                      "funções do RuFaS a partir desta planilha e do <code>NASEM_Comp.csv</code>.", styles['TableCell']),
            Paragraph("Função objetivo, variáveis de decisão e restrições NASEM.", styles['TableCell'])
        ],
        [
            Paragraph("<code>docs/<br/>OtimizadorNasem.pdf</code>", styles['TableCell']),
            Paragraph("Documento técnico do OtimizadorNasem, ferramenta independente (Simplex/HiGHS) que replica as "
                      "restrições NASEM do otimizador interno do RuFaS (SLSQP).", styles['TableCell']),
            Paragraph("Especificação das restrições e confirmação de que <code>rufas_id == feed_type</code>.", styles['TableCell'])
        ]
    ]
    t_sources = Table(sources_data, colWidths=[usable_width*0.33, usable_width*0.42, usable_width*0.25])
    t_sources.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1e293b')),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 4),
        ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#f8fafc')),
        ('BACKGROUND', (0,3), (-1,3), colors.HexColor('#f8fafc')),
        ('BACKGROUND', (0,5), (-1,5), colors.HexColor('#f8fafc')),
    ]))
    story.append(t_sources)
    story.append(Spacer(1, 14))

    # Section 5: Raw values comparison table
    story.append(Paragraph("5. Rastreabilidade Detalhada: Valores Brutos no NASEM_Comp.csv", styles['SectionHeader']))
    story.append(Paragraph(
        "Para fins de auditoria científica e validação das dietas, a tabela abaixo apresenta os valores brutos "
        "exatos (com precisão de 4 casas decimais) conforme registrados no arquivo <code>NASEM_Comp.csv</code>:",
        styles['BodyP']
    ))

    raw_headers = [
        Paragraph("<b>rufas_id</b>", styles['TableHead']),
        Paragraph("<b>Nome Oficial no NASEM (Name)</b>", styles['TableHead']),
        Paragraph("<b>DM bruto</b>", styles['TableHead']),
        Paragraph("<b>CP bruto</b>", styles['TableHead']),
        Paragraph("<b>NDF bruto</b>", styles['TableHead']),
        Paragraph("<b>Ca bruto</b>", styles['TableHead']),
        Paragraph("<b>P bruto</b>", styles['TableHead']),
    ]

    raw_rows = [raw_headers]
    raw_records = [
        (18, 'Bermudagrass hay', '93.3730', '10.9910', '65.4080', '0.4910', '0.1950'),
        (23, 'Blood meal, low dRUP', '90.8640', '95.0600', '0.0000', '0.1350', '0.2760'),
        (44, 'Corn grain dry, medium grind', '86.9070', '8.5140', '9.7610', '0.0350', '0.3090'),
        (51, 'Corn silage, typical', '35.3610', '7.7070', '40.9310', '0.2370', '0.2320'),
        (94, 'Grass lg mixt, grass hay, mid', '89.2390', '15.6080', '54.6850', '0.6250', '0.3290'),
        (100, 'Grass lg mixt, mix silage', '40.0210', '17.6790', '51.2320', '0.8700', '0.3410'),
        (110, 'Legume silage, mid-maturity', '42.8830', '20.4710', '43.1950', '1.2540', '0.3510'),
        (176, 'Sudangrass silage, mature', '31.4710', '9.5380', '66.6170', '0.4460', '0.2660'),
        (202, 'Whole milk', '12.5000', '25.4000', '0.0000', '1.0000', '0.7500'),
        (216, 'Calf starter 18CP high starch (Fed to calves)', '88.0000', '20.6000', '17.3700', '0.7500', '0.3700'),
        (301, 'Farm ES Mineral Mix', '99.4594', '0.0351', '0.1449', '15.4300', '2.3200'),
        (303, 'Farm ES Northeast BP Blend', '74.1127', '33.3732', '23.8769', '0.4217', '0.7297'),
    ]

    for rec in raw_records:
        raw_rows.append([
            Paragraph(str(rec[0]), styles['TableCellCenter']),
            Paragraph(rec[1], styles['TableCell']),
            Paragraph(rec[2], styles['TableCellRight']),
            Paragraph(rec[3], styles['TableCellRight']),
            Paragraph(rec[4], styles['TableCellRight']),
            Paragraph(rec[5], styles['TableCellRight']),
            Paragraph(rec[6], styles['TableCellRight']),
        ])

    t_raw = Table(raw_rows, colWidths=[45, 185, 62, 62, 62, 62, 62])
    t_raw_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#334155')),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('PADDING', (0,0), (-1,-1), 3),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#334155')),
    ]
    for r_idx in range(1, len(raw_rows)):
        if r_idx % 2 == 0:
            t_raw_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#f8fafc')))
        else:
            t_raw_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.white))
    t_raw.setStyle(TableStyle(t_raw_style))
    story.append(t_raw)
    story.append(Spacer(1, 14))

    # Section 6: Conclusion / Sign-off
    story.append(Paragraph("6. Validação e Conclusão", styles['SectionHeader']))
    story.append(Paragraph(
        "A composição dos 12 alimentos confere com o <code>NASEM_Comp.csv</code> usado pelo RuFaS, e todos os itens "
        "têm preço com fonte documentada. A planilha fornece a <b>função objetivo</b> (custos em R$/kg MS) do modelo de "
        "programação linear. As restrições nutricionais dependem também de colunas do <code>NASEM_Comp.csv</code> "
        "(energia, proteína metabolizável, extrato etéreo) e das exigências dos animais simuladas no RuFaS; o modelo "
        "completo é gerado por <code>otimizacao/gerar_modelo_pl.py</code>.",
        styles['BodyP']
    ))
    story.append(Paragraph(
        "<b>Limitações dos preços:</b> só milho, leite e silagem de milho têm cotação oficial (CEPEA e CNA/Cepea). "
        "As silagens 100, 110 e 176 e o blend 303 foram precificados por aproximação, e os anúncios usados não "
        "incluem frete até Goiás. Esses itens são os primeiros candidatos à análise de sensibilidade econômica.",
        styles['BodyP']
    ))

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    
    # Copy to docs/ as well
    import shutil
    shutil.copy2(pdf_filename, docs_pdf_filename)
    print('Generated PDF at:', pdf_filename)
    print('Copied PDF to:', docs_pdf_filename)

if __name__ == '__main__':
    generate_pdf()
